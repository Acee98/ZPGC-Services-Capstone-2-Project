"""
Convert ZPGC docs/*.md → docs/docx/*.docx
Usage:
  python tools/md_to_docx.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor


DOCS = Path(r"C:\xampp\htdocs\CP2_V1.5\docs")
OUT = DOCS / "docx"
SKIP = {"BACKLOG.corrupt.bak.md"}


def clean_text(s: str) -> str:
    """Remove XML-illegal controls; keep tabs/newlines."""
    if s is None:
        return ""
    s = s.replace("\x00", "")
    # Replace common arrows that confuse some Windows consoles when printing
    # but keep them in the document as ASCII equivalents for safety.
    replacements = {
        "\u2192": "->",
        "\u2190": "<-",
        "\u2014": "-",
        "\u2013": "-",
        "\u2026": "...",
        "\u00a0": " ",
        "\ufeff": "",
    }
    for a, b in replacements.items():
        s = s.replace(a, b)
    out = []
    for ch in s:
        o = ord(ch)
        if ch in "\t\n\r":
            out.append(ch)
        elif o < 32 or (0x7F <= o <= 0x9F):
            continue
        else:
            out.append(ch)
    return "".join(out)


def set_code_style(run):
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")


def add_horizontal_line(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    bottom = OxmlElement("w:pBdr")
    border = OxmlElement("w:bottom")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "6")
    border.set(qn("w:space"), "1")
    border.set(qn("w:color"), "888888")
    bottom.append(border)
    p._p.get_or_add_pPr().append(bottom)


def add_runs_with_inline(paragraph, text: str):
    """Support **bold**, `code`, and plain text (no nested mixing extremes)."""
    text = clean_text(text)
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        token = m.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            set_code_style(run)
            run.font.color.rgb = RGBColor(0x61, 0x01, 0x07)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.rows[i].cells[j]
            cell.text = clean_text(row[j] if j < len(row) else "")
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    if i == 0:
                        run.bold = True
    doc.add_paragraph()


def parse_md_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        line = lines[i].strip()
        # skip separator |---|---|
        if re.match(r"^\|[\s:\-|]+\|$", line.replace(" ", "")) or re.match(
            r"^\|[\s\-:|]+\|$", line
        ):
            i += 1
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)
        i += 1
    return rows, i


def convert_file(md_path: Path, out_path: Path):
    text = md_path.read_text(encoding="utf-8", errors="replace")
    text = clean_text(text.replace("\r\n", "\n"))
    lines = text.split("\n")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tr = title.add_run(clean_text(md_path.stem.replace("_", " ")))
    tr.bold = True
    tr.font.size = Pt(18)
    tr.font.color.rgb = RGBColor(0x61, 0x01, 0x07)
    doc.add_paragraph(f"Source: {md_path.name}")
    add_horizontal_line(doc)

    i = 0
    in_code = False
    code_lang = ""
    code_buf: list[str] = []
    in_mermaid = False

    while i < len(lines):
        line = lines[i]

        # fenced code
        fence = re.match(r"^```(\w*)\s*$", line)
        if fence:
            if not in_code:
                in_code = True
                code_lang = (fence.group(1) or "").lower()
                code_buf = []
                in_mermaid = code_lang == "mermaid"
            else:
                # close fence
                body = clean_text("\n".join(code_buf))
                if in_mermaid:
                    p = doc.add_paragraph()
                    r = p.add_run("[Mermaid diagram - see source .md]")
                    r.italic = True
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
                else:
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Inches(0.15)
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(4)
                    run = p.add_run(body if body else " ")
                    set_code_style(run)
                    if code_lang:
                        label = doc.add_paragraph()
                        lr = label.add_run(f"code ({code_lang})")
                        lr.italic = True
                        lr.font.size = Pt(8)
                        lr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
                in_code = False
                code_buf = []
                in_mermaid = False
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # blank
        if not line.strip():
            i += 1
            continue

        # horizontal rule
        if re.match(r"^---+\s*$", line.strip()):
            add_horizontal_line(doc)
            i += 1
            continue

        # image ![alt](path)
        img = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", line.strip())
        if img:
            alt, rel = img.group(1), img.group(2)
            img_path = (md_path.parent / rel).resolve()
            cap = doc.add_paragraph()
            cr = cap.add_run(clean_text(alt or rel))
            cr.italic = True
            cr.font.size = Pt(9)
            if img_path.is_file():
                try:
                    doc.add_picture(str(img_path), width=Inches(5.8))
                except Exception:
                    p = doc.add_paragraph()
                    p.add_run(f"[Image could not be embedded: {rel}]").italic = True
            else:
                p = doc.add_paragraph()
                p.add_run(f"[Image missing: {rel}]").italic = True
            i += 1
            continue

        # headings
        h = re.match(r"^(#{1,6})\s+(.*)$", line)
        if h:
            level = min(len(h.group(1)), 4)
            heading = doc.add_heading(clean_text(h.group(2).strip()), level=level)
            for run in heading.runs:
                if level == 1:
                    run.font.color.rgb = RGBColor(0x61, 0x01, 0x07)
            i += 1
            continue

        # table
        if line.strip().startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            rows, ni = parse_md_table(lines, i)
            add_table(doc, rows)
            i = ni
            continue

        # bullet list
        if re.match(r"^[-*]\s+", line):
            text_item = re.sub(r"^[-*]\s+", "", line)
            p = doc.add_paragraph(style="List Bullet")
            add_runs_with_inline(p, text_item)
            i += 1
            continue

        # numbered list
        if re.match(r"^\d+\.\s+", line):
            text_item = re.sub(r"^\d+\.\s+", "", line)
            p = doc.add_paragraph(style="List Number")
            add_runs_with_inline(p, text_item)
            i += 1
            continue

        # blockquote
        if line.lstrip().startswith(">"):
            text_item = re.sub(r"^>\s?", "", line.lstrip())
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            run = p.add_run(clean_text(text_item))
            run.italic = True
            i += 1
            continue

        # normal paragraph
        p = doc.add_paragraph()
        add_runs_with_inline(p, line)
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"OK  {md_path.name} -> {out_path.name}")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    files = sorted(DOCS.glob("*.md"))
    converted = 0
    for md in files:
        if md.name in SKIP:
            print(f"SKIP {md.name}")
            continue
        out = OUT / (md.stem + ".docx")
        try:
            convert_file(md, out)
            converted += 1
        except Exception as exc:
            print(f"FAIL {md.name}: {exc}", file=sys.stderr)
    print(f"\nDone: {converted} file(s) in {OUT}")


if __name__ == "__main__":
    main()
